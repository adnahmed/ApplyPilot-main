"""Spreadsheet previews are inert; only an explicit commit creates targets."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from uuid import uuid4

from . import store
from applypilot.discovery.repository import transaction

ALIASES = {"company_name": "company", "name": "company", "company_website": "website",
           "url": "website", "careers": "careers_url", "job": "job_url", "linkedin": "profile_url",
           "contact": "contact_name", "role": "contact_role", "message": "instructions", "notes": "instructions"}


def parse(content, filename="paste.tsv"):
    if len(content) > 11_000_000:
        raise ValueError("Import exceeds the dashboard's 11 MB upload limit")
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            rows = list(book.active.iter_rows(values_only=True))
        finally:
            book.close()
    else:
        value = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(value[:8192], delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel_tab if "\t" in value else csv.excel
        rows = list(csv.reader(io.StringIO(value), dialect))
    rows = [list(row) for row in rows if any(value is not None and str(value).strip() for value in row)]
    if not rows:
        raise ValueError("The spreadsheet is empty")
    headers = [str(x or "").strip() for x in rows[0]]
    if len(set(headers)) != len(headers) or any(not x for x in headers):
        raise ValueError("Use unique, nonempty column headings")
    return {"headers": headers, "rows": [[str(x) if x is not None else "" for x in row] for row in rows[1:]],
            "mapping": {h: ALIASES.get(re.sub(r"[\s-]+", "_", h.lower()), re.sub(r"[\s-]+", "_", h.lower()))
                        if ALIASES.get(re.sub(r"[\s-]+", "_", h.lower()), re.sub(r"[\s-]+", "_", h.lower())) in store.FIELDS else ""
                        for h in headers}}


def preview(conn, parsed, mapping=None, defaults=None):
    mapping = parsed["mapping"] if mapping is None else mapping
    defaults = defaults or {}
    if set(defaults) - set(store.FIELDS) or any(v and v not in store.FIELDS for v in mapping.values()):
        raise ValueError("Unknown import field")
    selected_fields = [v for v in mapping.values() if v]
    if len(set(selected_fields)) != len(selected_fields):
        raise ValueError("Map each target field only once")
    rows, seen = [], {}
    for index, cells in enumerate(parsed["rows"]):
        data = dict(defaults)
        for column, header in enumerate(parsed["headers"]):
            field = mapping.get(header)
            if field and column < len(cells) and str(cells[column]).strip():
                data[field] = str(cells[column]).strip()
        item = {"index": index, "row": index + 2, "data": data, "error": "", "duplicate": None}
        try:
            item["data"] = value = store.validate(data)
            key = store.identity(value)
            match = conn.execute("SELECT id,company FROM company_targets WHERE identity_key=?", (key,)).fetchone()
            if match:
                item["duplicate"] = {"id": match[0], "company": match[1]}
            elif key in seen:
                item["duplicate"] = {"row": seen[key]}
            seen[key] = index + 2
        except ValueError as exc:
            item["error"] = str(exc)
        rows.append(item)
    digest = hashlib.sha256(store.dump(rows).encode()).hexdigest()
    with transaction(conn, resource="company-targets"):
        conn.execute("INSERT INTO company_target_imports(id,digest,rows_json,created_at) VALUES(?,?,?,?) ON CONFLICT(digest) DO NOTHING",
                     (uuid4().hex, digest, store.dump(rows), store.now()))
        record = conn.execute("SELECT id,committed_at FROM company_target_imports WHERE digest=?", (digest,)).fetchone()
    return {"id": record[0], "committed_at": record[1], "rows": rows, "mapping": mapping, "parsed": parsed}


def commit(conn, import_id, selected):
    if not isinstance(selected, list) or any(type(x) is not int for x in selected):
        raise ValueError("Select the rows to import")
    with transaction(conn, resource="company-targets"):
        record = conn.execute("SELECT * FROM company_target_imports WHERE id=? FOR UPDATE", (import_id,)).fetchone()
        if not record:
            raise LookupError("Import preview not found")
        if record["committed_at"]:
            return json.loads(record["result_json"])
        rows = json.loads(record["rows_json"])
        if not selected or any(x < 0 or x >= len(rows) for x in selected):
            raise ValueError("Select at least one valid preview row")
        chosen = [rows[i] for i in sorted(set(selected))]
        if any(x["error"] for x in chosen):
            raise ValueError("Correct invalid rows before importing")
        results = [store.save(conn, row["data"]) for row in chosen]
        result = {"created": sum(not x["duplicate"] for x in results),
                  "duplicates": sum(x["duplicate"] for x in results), "ids": [x["target"]["id"] for x in results]}
        conn.execute("UPDATE company_target_imports SET committed_at=?,result_json=? WHERE id=?",
                     (store.now(), store.dump(result), import_id))
        return result
