"""Owner-selected document intake; bounded text only, never executes attachments."""
from __future__ import annotations
import base64
import io
import json
from pathlib import Path
import re
import zipfile
import yaml
from .storage import atomic, digest, now, public, read_json, setup_lock, write_json

MAX_FILE = 8_000_000
MAX_TEXT = 120_000


def clean_text(text):
    # Explicit secret-labelled lines are not useful candidate evidence.
    lines = [line for line in str(text).replace("\x00", "").splitlines()
             if not re.match(r"\s*(?:password|api[_ -]?key|access[_ -]?token|client[_ -]?secret|refresh[_ -]?token)\s*[:=]", line, re.I)]
    return "\n".join(lines)[:MAX_TEXT].strip()


def extract(name, raw):
    suffix = Path(name).suffix.lower()
    if not raw or len(raw) > MAX_FILE:
        raise ValueError("Choose a nonempty document of at most 8 MB")
    payload = None
    if suffix in {".txt", ".md", ".json", ".yaml", ".yml"}:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeError:
            raise ValueError("Text files must be UTF-8; use PDF or DOCX for other formats") from None
        if suffix == ".json":
            payload = json.loads(text)
            if not isinstance(payload, dict) or "personal" not in payload:
                raise ValueError("JSON intake accepts a profile.json mapping with a personal section; use the Gmail credential control for OAuth JSON")
            payload = public(payload)
            text = json.dumps(payload, ensure_ascii=False, indent=2)
        elif suffix in {".yaml", ".yml"}:
            payload = yaml.safe_load(text)
            if not isinstance(payload, dict):
                raise ValueError("Search YAML must be a mapping")
            from applypilot.discovery.intent import compile_search_intent
            compile_search_intent(payload)
    elif suffix == ".pdf":
        if not raw.startswith(b"%PDF-"):
            raise ValueError("The file does not contain a PDF header")
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw), strict=False)
        if reader.is_encrypted:
            raise ValueError("Unlock the PDF locally before importing it")
        if len(reader.pages) > 40:
            raise ValueError("Use a résumé or excerpt with at most 40 PDF pages")
        parts = []
        length = 0
        for page in reader.pages:
            part = page.extract_text() or ""
            parts.append(part)
            length += len(part)
            if length > MAX_TEXT: break
        text = "\n\n".join(parts)
    elif suffix == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(x.file_size for x in infos) > 25_000_000:
                    raise ValueError("DOCX expanded size exceeds the intake limit")
                if any("vbaProject" in x.filename for x in infos):
                    raise ValueError("Macro-enabled documents are not accepted")
            from docx import Document
            document = Document(io.BytesIO(raw))
            text = "\n".join([p.text for p in document.paragraphs] +
                             [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        except (zipfile.BadZipFile, KeyError):
            raise ValueError("The file is not a readable DOCX document") from None
    else:
        raise ValueError("Use PDF, DOCX, TXT, Markdown, profile.json, or searches.yaml")
    text = clean_text(text)
    if len(text) < 30:
        raise ValueError("Too little selectable text was found. For scans, paste a text version; no automatic OCR or external upload is performed.")
    return text, payload, suffix


def intake(root, data, *, progress=lambda **kwargs: None):
    name = str(data.get("name") or "notes.txt").replace("\\", "/").split("/")[-1][:120]
    kind = data.get("kind", "supporting")
    if kind not in {"resume", "supporting", "profile", "searches"}: raise ValueError("Unknown intake kind")
    if str(data.get('text') or '').strip() and data.get('content_base64'):
        raise ValueError('Choose one document or paste its text for this import')
    if str(data.get('text') or '').strip():
        raw = str(data["text"]).encode()
        name = "pasted-notes.txt" if kind != "resume" else "pasted-resume.txt"
    else:
        try:
            raw = base64.b64decode(data.get("content_base64", ""), validate=True)
        except (ValueError, TypeError):
            raise ValueError("Invalid document encoding") from None
    text, payload, suffix = extract(name, raw)
    if suffix == ".json": kind = "profile"
    if suffix in {".yaml", ".yml"}: kind = "searches"
    identifier = digest(raw)[:20]
    record = {"id": identifier,"name": name,"kind": kind,"sha256":digest(raw),"characters":len(text),
              "text":text,"payload":payload,"extension":suffix,"imported_at":now(),
              "truncated":len(text)>=MAX_TEXT,"model_shared_by_intake":False}
    with setup_lock(root):
        path = root/"setup"/"sources.json"
        sources = read_json(path, [])
        sources = [r for r in sources if r["id"] != identifier]
        if len(sources)>=20 or sum(r["characters"] for r in sources)+len(text)>500_000:
            raise ValueError("Intake is limited to 20 documents / 500,000 text characters; remove old sources first")
        # Only originals potentially used as the master résumé are retained. Credentials never enter intake.
        if kind=="resume" and suffix==".pdf":
            atomic(root/"setup"/"uploads"/(identifier+".pdf"),raw)
        sources.append(record)
        write_json(path,sources)
    # Supported missing facts become available to the profile immediately.
    from .planning import load_draft, propose
    try:
        proposed=propose(root,{"revision":load_draft(root)["revision"],"use_ai":False,"apply_profile":True}, progress=progress)
        prepared=proposed["draft_saved"]
        notice="Material saved. Supported missing profile fields were filled; existing answers were preserved."
    except ValueError:
        prepared=False
        notice="Material saved. Profile completion needs attention; use Fill from saved material to retry."
    return {"draft_prepared":prepared,"source":{k:v for k,v in record.items() if k not in {"text","payload"}},"note":notice}
